"""Import and display Feedlync Loaded Mixes → By Ingredient usage."""

from __future__ import annotations

import datetime as dt
import io
import threading
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.models import HERD_FARM_OPTIONS, FeedUsageRecord
from app.services.farm_schedule import FARM_LABELS, normalize_farm
from app.services.feed_usage_settings import (
    assigned_ration_names,
    ingredient_inclusion_lookup,
    is_usage_ingredient_included,
    seed_ration_assignments_if_empty,
)
from app.services.feedlync_api import (
    USAGE_RATIONS_BY_FARM,
    fetch_loaded_mix_ingredient_usage,
    month_bounds,
    previous_calendar_month,
)
from app.services.feedlync_auth import FeedlyncAuthError

XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
_USAGE_XLSX_HEADERS = ("Ingredient", "As Fed (kg)", "As Fed (MT)")
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_KG_NUMBER_FORMAT = "#,##0"
_MT_NUMBER_FORMAT = "#,##0.0000"

_lock = threading.Lock()
_import_status: dict[str, Any] = {
    "status": "idle",
    "message": "",
    "latest_import": None,
    "rows_imported": 0,
    "month": None,
    "needs_auth": False,
}


def get_import_status() -> dict[str, Any]:
    with _lock:
        return dict(_import_status)


def _set_status(**kwargs: Any) -> None:
    with _lock:
        _import_status.update(kwargs)


def is_import_running() -> bool:
    with _lock:
        return _import_status.get("status") == "running"


def mark_import_started(month: str) -> None:
    _set_status(
        status="running",
        message="Starting Feedlync usage import…",
        rows_imported=0,
        month=month,
        needs_auth=False,
    )


def resolve_usage_month(value: str | None, *, today: dt.date | None = None) -> dt.date:
    if value:
        text = value.strip()
        if len(text) == 7 and text[4] == "-":
            text = text + "-01"
        try:
            parsed = dt.date.fromisoformat(text)
        except ValueError as exc:
            raise ValueError(f"Invalid month: {value}") from exc
        return parsed.replace(day=1)
    return previous_calendar_month(today)


def _round_kg(value: float) -> float:
    return round(value, 2)


def _round_money(value: float) -> float:
    return round(value, 2)


def build_usage_report(
    rows: list[dict[str, Any]],
    *,
    period_start: dt.date,
    period_end: dt.date,
    farm: str,
    rations: list[str] | None = None,
    import_timestamp: dt.datetime | None = None,
) -> dict[str, Any]:
    farm_key = normalize_farm(farm)
    ration_names = (
        list(rations)
        if rations is not None
        else list(USAGE_RATIONS_BY_FARM.get(farm_key, ()))
    )
    days = (period_end - period_start).days + 1
    ingredients = sorted(
        rows,
        key=lambda row: float(row.get("as_fed_kg") or 0),
        reverse=True,
    )
    total_as_fed = sum(float(row.get("as_fed_kg") or 0) for row in ingredients)
    total_dm = sum(float(row.get("dm_kg") or 0) for row in ingredients)
    total_cost = sum(float(row.get("cost") or 0) for row in ingredients)
    return {
        "month": period_start.strftime("%Y-%m"),
        "period_start": period_start.isoformat(),
        "period_end": period_end.isoformat(),
        "days": days,
        "farm": farm_key,
        "farm_label": FARM_LABELS[farm_key],
        "rations": ration_names,
        "source": "Loaded Mixes → By Ingredient",
        "ingredients": [
            {
                "ingredient_name": row.get("ingredient_name") or "",
                "as_fed_kg": _round_kg(float(row.get("as_fed_kg") or 0)),
                "dm_kg": _round_kg(float(row.get("dm_kg") or 0)),
                "cost": _round_money(float(row.get("cost") or 0)),
            }
            for row in ingredients
        ],
        "totals": {
            "as_fed_kg": _round_kg(total_as_fed),
            "dm_kg": _round_kg(total_dm),
            "cost": _round_money(total_cost),
        },
        "averages": {
            "as_fed_kg": _round_kg(total_as_fed / days) if days else 0.0,
            "dm_kg": _round_kg(total_dm / days) if days else 0.0,
            "cost": _round_money(total_cost / days) if days else 0.0,
        },
        "latest_import": import_timestamp.isoformat() if import_timestamp else None,
        "row_count": len(ingredients),
    }


def _as_fed_mt(kg: float) -> float:
    return round(float(kg or 0) / 1000, 4)


def _xlsx_display_width(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, bool):
        return len(str(value))
    if isinstance(value, int):
        return len(f"{value:,}")
    if isinstance(value, float):
        return len(f"{value:,.4f}".rstrip("0").rstrip("."))
    return len(str(value))


def _fit_xlsx_columns(ws) -> None:
    for index in range(1, (ws.max_column or 1) + 1):
        letter = get_column_letter(index)
        widest = 0
        for cell in ws[letter]:
            widest = max(widest, _xlsx_display_width(cell.value))
        ws.column_dimensions[letter].width = min(max(widest + 2, 12), 48)


def _border_xlsx_cells(ws) -> None:
    for row in ws.iter_rows(
        min_row=1,
        max_row=max(ws.max_row, 1),
        min_col=1,
        max_col=max(ws.max_column, 1),
    ):
        for cell in row:
            cell.border = _THIN_BORDER


def build_usage_xlsx(report: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    farm_label = str(report.get("farm_label") or report.get("farm") or "Feed Usage")
    ws.title = farm_label[:31]
    header_font = Font(bold=True)
    total_font = Font(bold=True)
    ws.append(list(_USAGE_XLSX_HEADERS))
    for cell in ws[1]:
        cell.font = header_font
    ws["B1"].alignment = Alignment(horizontal="right")
    ws["C1"].alignment = Alignment(horizontal="right")

    for row in report.get("ingredients") or []:
        kg = float(row.get("as_fed_kg") or 0)
        ws.append([row.get("ingredient_name") or "", kg, _as_fed_mt(kg)])
        data_row = ws[ws.max_row]
        data_row[1].number_format = _KG_NUMBER_FORMAT
        data_row[2].number_format = _MT_NUMBER_FORMAT
        data_row[1].alignment = Alignment(horizontal="right", vertical="center")
        data_row[2].alignment = Alignment(horizontal="right", vertical="center")

    totals = report.get("totals") or {}
    total_kg = float(totals.get("as_fed_kg") or 0)
    ws.append(["Total", total_kg, _as_fed_mt(total_kg)])
    total_row = ws[ws.max_row]
    for cell in total_row:
        cell.font = total_font
    total_row[1].number_format = _KG_NUMBER_FORMAT
    total_row[2].number_format = _MT_NUMBER_FORMAT
    total_row[1].alignment = Alignment(horizontal="right", vertical="center")
    total_row[2].alignment = Alignment(horizontal="right", vertical="center")

    _fit_xlsx_columns(ws)
    _border_xlsx_cells(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def get_usage_report(db: Session, *, month: dt.date, farm: str) -> dict[str, Any]:
    farm_key = normalize_farm(farm)
    period_start, period_end = month_bounds(month.year, month.month)
    records = list(
        db.scalars(
            select(FeedUsageRecord)
            .where(
                FeedUsageRecord.period_start == period_start,
                FeedUsageRecord.farm == farm_key,
            )
            .order_by(FeedUsageRecord.as_fed_kg.desc(), FeedUsageRecord.ingredient_name)
        ).all()
    )
    latest_import = None
    if records:
        latest_import = max(
            (record.import_timestamp for record in records if record.import_timestamp),
            default=None,
        )
    seed_ration_assignments_if_empty(db)
    inclusion = ingredient_inclusion_lookup(db)
    return build_usage_report(
        [
            record.to_dict()
            for record in records
            if is_usage_ingredient_included(record.ingredient_name, inclusion)
        ],
        period_start=period_start,
        period_end=period_end,
        farm=farm_key,
        rations=assigned_ration_names(db, farm_key),
        import_timestamp=latest_import,
    )


def import_feed_usage(db: Session, *, month: dt.date) -> dict[str, Any]:
    period_start, period_end = month_bounds(month.year, month.month)
    month_key = period_start.strftime("%Y-%m")
    _set_status(
        status="running",
        message=f"Fetching Loaded Mixes for {month_key}…",
        rows_imported=0,
        month=month_key,
        needs_auth=False,
    )
    try:
        rows = fetch_loaded_mix_ingredient_usage(
            db, period_start=period_start, period_end=period_end
        )
        if not rows:
            raise ValueError(
                f"No Loaded Mixes ingredient rows returned from Feedlync for {month_key}."
            )

        import_ts = dt.datetime.now()
        db.execute(
            delete(FeedUsageRecord).where(FeedUsageRecord.period_start == period_start)
        )
        db.flush()

        for row in rows:
            farm = str(row.get("farm") or "").strip().upper()
            if farm not in HERD_FARM_OPTIONS:
                continue
            db.add(
                FeedUsageRecord(
                    period_start=period_start,
                    period_end=period_end,
                    farm=farm,
                    ingredient_name=str(row["ingredient_name"]),
                    as_fed_kg=float(row["as_fed_kg"]),
                    dm_kg=float(row["dm_kg"]),
                    cost=float(row["cost"]),
                    import_timestamp=import_ts,
                )
            )

        db.commit()
        latest_import = import_ts.isoformat()
        result = {
            "rows_imported": len(rows),
            "month": month_key,
            "period_start": period_start.isoformat(),
            "period_end": period_end.isoformat(),
            "latest_import": latest_import,
        }
        _set_status(
            status="complete",
            message=f"Imported {len(rows)} ingredients for {month_key}.",
            latest_import=latest_import,
            rows_imported=len(rows),
            month=month_key,
            needs_auth=False,
        )
        return result
    except FeedlyncAuthError as exc:
        db.rollback()
        _set_status(status="error", message=str(exc), month=month_key, needs_auth=True)
        raise
    except Exception as exc:
        db.rollback()
        _set_status(status="error", message=str(exc), month=month_key, needs_auth=False)
        raise


def run_usage_import_in_background(db_factory, month: dt.date) -> None:
    db = db_factory()
    try:
        import_feed_usage(db, month=month)
    except Exception:
        pass
    finally:
        db.close()
