"""Combined CM and GAD inventory table for Genetics — Animal Data."""

from __future__ import annotations

import csv
import datetime as dt
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.models import GenomicResult, HerdInventory
from app.services.events_common import normalize_farms
from app.services.genomic_import import normalize_hbn

YS_DCC_LIMIT = 100
MALE_RC = 8
MAX_INCLUDED_CBRD = 101
BULL_RPRO = "BULL"

FARM_GROUPS: tuple[str, ...] = ("CM Cows", "CM YS", "GAD Cows", "GAD YS")

XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

COLUMNS: tuple[tuple[str, str], ...] = (
    ("id", "ID"),
    ("etag", "ETAG"),
    ("bdat", "BDAT"),
    ("lact", "LACT"),
    ("cci", "CCI"),
    ("dreg", "DREG"),
    ("sreg", "SREG"),
    ("mgreg", "MGREG"),
    ("ggreg", "GGREG"),
    ("fdat", "FDAT"),
    ("due", "DUE"),
    ("hdat", "HDAT"),
    ("rpro", "RPRO"),
    ("cbrd", "CBRD"),
    ("dim", "DIM"),
    ("tbrd", "TBRD"),
    ("dcc", "DCC"),
    ("farm_group", "Farm Group"),
)

_DATE_KEYS = {"bdat", "fdat", "due", "hdat"}
_XLSX_WIDTHS = (
    12, 20, 12, 8, 10, 18, 18, 18, 18, 12, 12, 12, 10, 8, 8, 8, 8, 14,
)


def farm_group_label(farm: str | None, lact: Any, dcc: Any) -> str:
    """Lact 0 and DCC under 100 are youngstock; every other animal is a cow."""
    prefix = (farm or "").strip().upper()
    kind = "YS" if _is_youngstock(lact, dcc) else "Cows"
    return f"{prefix} {kind}"


def normalize_farm_groups(groups: list[str] | None) -> list[str]:
    if not groups:
        return list(FARM_GROUPS)
    return [group for group in groups if group in FARM_GROUPS]


def list_animal_data(
    db: Session,
    *,
    farms: list[str] | None = None,
    groups: list[str] | None = None,
) -> dict[str, Any]:
    """Dairy animals from both herd files, excluding bulls, males, and beef.

    Excludes RPRO BULL, RC 8, and CBRD above 101. Farm Group is CM/GAD Cows
    or CM/GAD YS from the source file, lactation, and days carried calf.
    CCI comes from genomic results, matched on ear-tag digits.
    """
    selected_farms = normalize_farms(farms)
    selected_groups = normalize_farm_groups(groups)
    if not selected_farms or not selected_groups:
        return {"rows": [], "total": 0}

    query = (
        select(
            HerdInventory.farm,
            HerdInventory.cow_id,
            HerdInventory.etag,
            HerdInventory.bdat,
            HerdInventory.lact,
            HerdInventory.dreg,
            HerdInventory.sreg,
            HerdInventory.mgreg,
            HerdInventory.ggreg,
            HerdInventory.fdat,
            HerdInventory.due,
            HerdInventory.hdat,
            HerdInventory.rpro,
            HerdInventory.cbrd,
            HerdInventory.dim,
            HerdInventory.tbrd,
            HerdInventory.dcc,
        )
        .where(HerdInventory.farm.in_(selected_farms))
        .where(or_(HerdInventory.rc.is_(None), HerdInventory.rc != MALE_RC))
        .where(
            or_(
                HerdInventory.cbrd.is_(None),
                HerdInventory.cbrd <= MAX_INCLUDED_CBRD,
            )
        )
        .where(
            or_(
                HerdInventory.rpro.is_(None),
                func.upper(func.trim(HerdInventory.rpro)) != BULL_RPRO,
            )
        )
    )

    cci_by_tag = _genomic_cci_by_ear_tag(db)

    rows: list[dict[str, Any]] = []
    for record in db.execute(query).all():
        (
            farm,
            cow_id,
            etag,
            bdat,
            lact,
            dreg,
            sreg,
            mgreg,
            ggreg,
            fdat,
            due,
            hdat,
            rpro,
            cbrd,
            dim,
            tbrd,
            dcc,
        ) = record
        group = farm_group_label(farm, lact, dcc)
        if group not in selected_groups:
            continue
        rows.append(
            {
                "id": (cow_id or "").strip(),
                "etag": (etag or "").strip(),
                "bdat": _date_text(bdat),
                "lact": _number(lact),
                "cci": _cci_for_etag(cci_by_tag, etag),
                "dreg": (dreg or "").strip(),
                "sreg": (sreg or "").strip(),
                "mgreg": (mgreg or "").strip(),
                "ggreg": (ggreg or "").strip(),
                "fdat": _date_text(fdat),
                "due": _date_text(due),
                "hdat": _date_text(hdat),
                "rpro": (rpro or "").strip(),
                "cbrd": _number(cbrd),
                "dim": _number(dim),
                "tbrd": _number(tbrd),
                "dcc": _number(dcc),
                "farm_group": group,
            }
        )

    rows.sort(key=lambda row: (row["farm_group"], _id_sort_key(row["id"]), row["etag"]))
    return {"rows": rows, "total": len(rows)}


def build_animal_data_csv(rows: list[dict[str, Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([label for _key, label in COLUMNS])
    for row in rows:
        writer.writerow([_csv_value(row.get(key)) for key, _label in COLUMNS])
    return buffer.getvalue().encode("utf-8-sig")


def build_animal_data_xlsx(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Animal Data"
    ws.append([label for _key, label in COLUMNS])
    for row in rows:
        ws.append([_xlsx_value(key, row.get(key)) for key, _label in COLUMNS])

    for index, width in enumerate(_XLSX_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _genomic_cci_by_ear_tag(db: Session) -> dict[str, int | float]:
    """Map ear-tag digits to genomic CCI. Ear tag is preferred over HBN."""
    lookup: dict[str, int | float] = {}
    rows = db.execute(
        select(GenomicResult.eartag, GenomicResult.hbn, GenomicResult.cci)
    ).all()
    for eartag, hbn, cci in rows:
        value = _number(cci)
        if value is None:
            continue
        for raw in (eartag, hbn):
            key = normalize_hbn(raw)
            if key and key not in lookup:
                lookup[key] = value
    return lookup


def _cci_for_etag(lookup: dict[str, int | float], etag: str | None) -> int | float | None:
    key = normalize_hbn(etag)
    if not key:
        return None
    return lookup.get(key)


def _is_youngstock(lact: Any, dcc: Any) -> bool:
    lact_value = _number(lact)
    dcc_value = _number(dcc)
    if lact_value is None or dcc_value is None:
        return False
    return lact_value == 0 and dcc_value < YS_DCC_LIMIT


def _number(value: Any) -> int | float | None:
    if value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number != number:  # NaN
        return None
    if number.is_integer():
        return int(number)
    return number


def _date_text(value: dt.date | None) -> str:
    if value is None:
        return ""
    return value.isoformat()


def _csv_value(value: Any) -> Any:
    return "" if value is None else value


def _xlsx_value(key: str, value: Any) -> Any:
    if value is None or value == "":
        return None
    if key in _DATE_KEYS and isinstance(value, str):
        try:
            return dt.date.fromisoformat(value)
        except ValueError:
            return value
    return value


def _id_sort_key(cow_id: str) -> tuple[int, int | str]:
    text = (cow_id or "").strip()
    if text.isdigit():
        return (0, int(text))
    return (1, text)
