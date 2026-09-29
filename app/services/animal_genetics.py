"""Animal Data: inventory identity plus genomic traits, matched on ear tag."""

from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import GenomicResult, HerdInventory
from app.services.custom_indexes import cm_index, load_index_settings, merge_index_settings
from app.services.events_common import normalize_farms
from app.services.genomic_import import normalize_hbn

XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

COLUMNS: tuple[tuple[str, str], ...] = (
    ("id", "ID"),
    ("etag", "ETAG"),
    ("sreg", "SREG"),
    ("lact", "LACT"),
    ("age", "Age"),
    ("cm", "£CM"),
    ("pli", "PLI"),
    ("milk_kg", "Milk Kg"),
    ("fat_kg", "Fat Kg"),
    ("protein_kg", "Protein kg"),
    ("fat_pct", "Fat %"),
    ("protein_pct", "Protein %"),
    ("fertility_index", "Fertility Index"),
    ("scc", "SCC"),
    ("life_span", "Lifespan"),
    ("mastitis", "Mastitis"),
    ("stature", "Stature"),
    ("chest_width", "Chest Width"),
)

EXPORT_COLUMNS: tuple[tuple[str, str], ...] = COLUMNS + (("gyon", "GYON"),)

_GENOMIC_FIELDS = (
    "pli",
    "milk_kg",
    "fat_kg",
    "protein_kg",
    "fat_pct",
    "protein_pct",
    "fertility_index",
    "scc",
    "life_span",
    "mastitis",
    "stature",
    "chest_width",
)

_XLSX_WIDTHS = (
    12, 20, 18, 8, 8, 10, 10, 12, 12, 14, 10, 12, 16, 10, 12, 12, 12, 14, 8,
)


def list_animal_genetics(
    db: Session,
    *,
    farms: list[str] | None = None,
) -> dict[str, Any]:
    """Inventory animals with genomic traits attached by ear-tag digits.

    ID, ETAG, LACT, and Age come from herd inventory. Age is months (days
    divided by 30). SREG, £CM, and the remaining traits come from genomic
    results. Only animals with a genomic result are included.
    """
    selected_farms = normalize_farms(farms)
    if not selected_farms:
        return {"rows": [], "total": 0}

    genomic_by_tag = _genomic_by_ear_tag(db)
    index_settings = merge_index_settings(load_index_settings(db))

    query = (
        select(
            HerdInventory.cow_id,
            HerdInventory.etag,
            HerdInventory.lact,
            HerdInventory.months_old,
            HerdInventory.aged,
        )
        .where(HerdInventory.farm.in_(selected_farms))
    )

    rows: list[dict[str, Any]] = []
    for cow_id, etag, lact, months_old, aged in db.execute(query).all():
        genomic = _genomic_for_etag(genomic_by_tag, etag)
        if genomic is None:
            continue
        row = {
            "id": (cow_id or "").strip(),
            "etag": (etag or "").strip(),
            "sreg": _display_sreg(genomic.sire_reg),
            "lact": _number(lact),
            "age": _age_months(months_old, aged),
            "cm": int(round(cm_index(genomic, index_settings, merged=True))),
        }
        for field in _GENOMIC_FIELDS:
            row[field] = _number(getattr(genomic, field, None))
        rows.append(row)

    rows.sort(key=lambda row: (_cm_rank(row), _id_sort_key(row["id"]), row["etag"]))
    return {"rows": rows, "total": len(rows)}


def build_animal_genetics_csv(rows: list[dict[str, Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([label for _key, label in EXPORT_COLUMNS])
    for row in rows:
        writer.writerow([_export_cell(key, row.get(key)) for key, _label in EXPORT_COLUMNS])
    return buffer.getvalue().encode("utf-8-sig")


def build_animal_genetics_xlsx(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Animal Data"
    ws.append([label for _key, label in EXPORT_COLUMNS])
    for row in rows:
        ws.append(
            [_export_cell(key, row.get(key), blank=None) for key, _label in EXPORT_COLUMNS]
        )
    for cell in ws["C"][1:]:
        cell.number_format = "@"

    for index, width in enumerate(_XLSX_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _display_sreg(value: str | None) -> str:
    """Unknown genomic sires are stored as UUUU… and shown as BULLHO."""
    text = (value or "").strip()
    compact = "".join(text.split()).upper()
    if compact.startswith("UUUU"):
        return "BULLHO"
    return text


def _export_sreg(value: Any) -> str:
    """Registration exports keep digits only. HO840… becomes 840…."""
    text = "" if value is None else str(value).strip()
    if text == "BULLHO":
        return text
    return "".join(character for character in text if character.isdigit())


def _export_cell(key: str, value: Any, *, blank: Any = "") -> Any:
    if key == "sreg":
        return _export_sreg(value)
    if key == "gyon":
        return 1
    if value is None:
        return blank
    return value


def _age_months(months_old: Any, aged: Any) -> int | None:
    """Inventory months, or days divided by 30 when months were not stored."""
    months = _number(months_old)
    if isinstance(months, int):
        return months
    if isinstance(months, float) and months.is_integer():
        return int(months)
    days = _number(aged)
    if days is None:
        return None
    return int(days) // 30


def _genomic_by_ear_tag(db: Session) -> dict[str, GenomicResult]:
    lookup: dict[str, GenomicResult] = {}
    for row in db.scalars(select(GenomicResult)).all():
        for raw in (row.eartag, row.hbn):
            key = normalize_hbn(raw)
            if key and key not in lookup:
                lookup[key] = row
    return lookup


def _genomic_for_etag(
    lookup: dict[str, GenomicResult], etag: str | None
) -> GenomicResult | None:
    key = normalize_hbn(etag)
    if not key:
        return None
    return lookup.get(key)


def _cm_rank(row: dict[str, Any]) -> tuple[int, float]:
    value = row.get("cm")
    if value is None:
        return (1, 0.0)
    return (0, -float(value))


def _id_sort_key(cow_id: str) -> tuple[int, int | str]:
    text = (cow_id or "").strip()
    if text.isdigit():
        return (0, int(text))
    return (1, text)


def _number(value: Any) -> int | float | None:
    if value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number != number:
        return None
    if number.is_integer():
        return int(number)
    return number
