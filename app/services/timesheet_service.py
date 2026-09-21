"""Staff time sheets: pay periods, hours, and holiday balances."""

from __future__ import annotations

import calendar
import datetime as dt
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    ACCOMMODATION_CADENCE_MONTHLY,
    ACCOMMODATION_CADENCE_WEEKLY,
    ACCOMMODATION_CADENCES,
    CM_TIMESHEET_PERIOD_START,
    DEFAULT_ANNUAL_LEAVE_DAYS,
    EMPLOYEE_STATUS_ARCHIVED,
    EMPLOYMENT_TYPE_EMPLOYED,
    EMPLOYMENT_TYPE_LABELS,
    EMPLOYMENT_TYPE_SELF_EMPLOYED,
    GAD_TIMESHEET_PERIOD_START,
    PAY_TYPE_HOURLY,
    PAY_TYPE_LABELS,
    PAY_TYPE_SALARY,
    TIMESHEET_FARM_BUSINESS,
    TIMESHEET_FARM_CM,
    TIMESHEET_FARM_LABELS,
    TIMESHEET_FARM_OPTIONS,
    TIMESHEET_FORTNIGHT_DAYS,
    Employee,
    EmployeeTimesheetEntry,
)
from app.services.crypto_fields import decrypt_field

WEEKS_PER_YEAR = 52
PERIOD_HORIZON_MONTHS = 12
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_HEADER_FILL = PatternFill("solid", fgColor="E7E6E6")
_TITLE_FILL = PatternFill("solid", fgColor="1E3A5F")
_SUBTITLE_FILL = PatternFill("solid", fgColor="E8EEF5")
_SALARY_FILL = PatternFill("solid", fgColor="D4EDDA")
_SELF_EMPLOYED_FILL = PatternFill("solid", fgColor="EFE6F7")
_TITLE_FONT = Font(bold=True, size=16, color="FFFFFF")
_SUBTITLE_FONT = Font(bold=True, size=11, color="1E3A5F")
_HEADER_FONT = Font(bold=True, size=10, color="1F1F1F")
_MONEY_FORMAT = "£#,##0.00"
_HOURS_FORMAT = "0.00"
_DATE_FORMAT = "DD/MM/YYYY"


class TimesheetError(Exception):
    """Time sheet operation failed."""


def normalize_farm(farm: str | None) -> str:
    value = (farm or "").strip().upper()
    if value not in TIMESHEET_FARM_OPTIONS:
        raise TimesheetError("Farm must be CM or GAD.")
    return value


def farm_label(farm: str) -> str:
    return TIMESHEET_FARM_LABELS[normalize_farm(farm)]


def farm_business(farm: str) -> str:
    return TIMESHEET_FARM_BUSINESS[normalize_farm(farm)]


def is_fortnightly(farm: str) -> bool:
    return normalize_farm(farm) == TIMESHEET_FARM_CM


def pay_cycle_start(farm: str) -> dt.date:
    return (
        CM_TIMESHEET_PERIOD_START
        if is_fortnightly(farm)
        else GAD_TIMESHEET_PERIOD_START
    )


def period_containing(farm: str, target: dt.date) -> tuple[dt.date, dt.date]:
    farm_key = normalize_farm(farm)
    start_anchor = pay_cycle_start(farm_key)
    if farm_key == TIMESHEET_FARM_CM:
        if target < start_anchor:
            start = start_anchor
        else:
            weeks = (target - start_anchor).days // TIMESHEET_FORTNIGHT_DAYS
            start = start_anchor + dt.timedelta(days=weeks * TIMESHEET_FORTNIGHT_DAYS)
        end = start + dt.timedelta(days=TIMESHEET_FORTNIGHT_DAYS - 1)
        return start, end

    if target < start_anchor:
        start = start_anchor
    else:
        start = dt.date(target.year, target.month, 1)
        if start < start_anchor:
            start = start_anchor
    last_day = calendar.monthrange(start.year, start.month)[1]
    return start, dt.date(start.year, start.month, last_day)


def resolve_period(farm: str, period_start: dt.date) -> tuple[dt.date, dt.date]:
    start, end = period_containing(farm, period_start)
    if start != period_start:
        raise TimesheetError(
            "period_start must be the first day of a pay period for this farm."
        )
    return start, end


def _add_months(value: dt.date, months: int) -> dt.date:
    month_index = value.year * 12 + (value.month - 1) + months
    year, month = divmod(month_index, 12)
    last_day = calendar.monthrange(year, month + 1)[1]
    return dt.date(year, month + 1, min(value.day, last_day))


def list_periods(
    farm: str,
    *,
    as_of: dt.date | None = None,
) -> list[dict[str, Any]]:
    farm_key = normalize_farm(farm)
    as_of = as_of or dt.date.today()
    cursor = pay_cycle_start(farm_key)
    horizon = _add_months(max(as_of, cursor), PERIOD_HORIZON_MONTHS)
    periods: list[dict[str, Any]] = []
    while cursor <= horizon:
        start, end = period_containing(farm_key, cursor)
        periods.append(serialize_period(farm_key, start, end))
        cursor = end + dt.timedelta(days=1)
    return periods


def current_period_start(farm: str, *, as_of: dt.date | None = None) -> dt.date:
    as_of = as_of or dt.date.today()
    start, _end = period_containing(farm, as_of)
    return start


def serialize_period(
    farm: str, start: dt.date, end: dt.date
) -> dict[str, Any]:
    farm_key = normalize_farm(farm)
    fortnightly = farm_key == TIMESHEET_FARM_CM
    if fortnightly:
        week_1_end = start + dt.timedelta(days=6)
        week_2_start = start + dt.timedelta(days=7)
        hours_columns = [
            {
                "key": "hours_week_1",
                "label": f"Hours {_format_range(start, week_1_end)}",
            },
            {
                "key": "hours_week_2",
                "label": f"Hours {_format_range(week_2_start, end)}",
            },
        ]
        cadence = "fortnightly"
        label = _format_range(start, end)
    else:
        hours_columns = [{"key": "hours_week_1", "label": "Hours"}]
        cadence = "monthly"
        label = start.strftime("%B %Y")
    return {
        "farm": farm_key,
        "start": start.isoformat(),
        "end": end.isoformat(),
        "label": label,
        "cadence": cadence,
        "hours_columns": hours_columns,
    }


def list_timesheet(
    db: Session,
    farm: str,
    *,
    period_start: dt.date | None = None,
    as_of: dt.date | None = None,
    include_rate: bool = False,
) -> dict[str, Any]:
    farm_key = normalize_farm(farm)
    as_of = as_of or dt.date.today()
    periods = list_periods(farm_key, as_of=as_of)
    if period_start is None:
        period_start = current_period_start(farm_key, as_of=as_of)
    start, end = resolve_period(farm_key, period_start)
    period = serialize_period(farm_key, start, end)
    if not any(item["start"] == period["start"] for item in periods):
        periods.append(period)
        periods.sort(key=lambda item: item["start"])

    staff = _active_staff_for_farm(db, farm_key)
    entries = {
        row.employee_id: row
        for row in db.scalars(
            select(EmployeeTimesheetEntry).where(
                EmployeeTimesheetEntry.farm == farm_key,
                EmployeeTimesheetEntry.period_start == start,
            )
        ).all()
    }
    holiday_hours_by_employee = _holiday_hours_by_employee(
        db, [row.id for row in staff]
    )
    return {
        "farm": farm_key,
        "farm_label": farm_label(farm_key),
        "business": farm_business(farm_key),
        "cadence": period["cadence"],
        "period": period,
        "periods": periods,
        "can_view_rate": include_rate,
        "rows": [
            _serialize_row(
                employee,
                entries.get(employee.id),
                holiday_hours_by_employee.get(employee.id, []),
                include_rate=include_rate,
                as_of=start,
                period_start=start,
                period_end=end,
            )
            for employee in staff
        ],
    }


def save_timesheet_row(
    db: Session,
    farm: str,
    payload: dict[str, Any],
    *,
    include_rate: bool = False,
) -> dict[str, Any]:
    farm_key = normalize_farm(farm)
    employee_id = int(payload["employee_id"])
    period_start = _parse_date(payload.get("period_start"), field="period_start")
    start, end = resolve_period(farm_key, period_start)
    employee = db.get(Employee, employee_id)
    if employee is None:
        raise TimesheetError("Employee not found.")
    if employee.status == EMPLOYEE_STATUS_ARCHIVED:
        raise TimesheetError("Archived staff cannot be added to a time sheet.")
    if (employee.business or "") != farm_business(farm_key):
        raise TimesheetError("Staff member does not belong to this farm.")

    entry = db.scalar(
        select(EmployeeTimesheetEntry).where(
            EmployeeTimesheetEntry.employee_id == employee_id,
            EmployeeTimesheetEntry.period_start == start,
        )
    )
    if entry is None:
        entry = EmployeeTimesheetEntry(
            employee_id=employee_id,
            farm=farm_key,
            period_start=start,
            period_end=end,
        )
        db.add(entry)

    entry.farm = farm_key
    entry.period_end = end
    is_salary = (employee.pay_type or PAY_TYPE_HOURLY) == PAY_TYPE_SALARY
    for key in (
        "hours_week_1",
        "holiday_hours",
        "dinner_break_hours",
        "mileage",
        "bonus",
        "loan_deduction",
        "other_deduction",
        "accommodation_deduction",
    ):
        if is_salary and key == "hours_week_1":
            continue
        if key in payload:
            setattr(entry, key, _optional_float(payload.get(key)))
    if is_fortnightly(farm_key):
        if not is_salary and "hours_week_2" in payload:
            entry.hours_week_2 = _optional_float(payload.get("hours_week_2"))
    else:
        entry.hours_week_2 = None
    if "other_remark" in payload:
        remark = payload.get("other_remark")
        entry.other_remark = (
            str(remark).strip() if remark is not None else ""
        ) or None

    if "holidays_carry_forward" in payload:
        employee.holidays_carry_forward = _optional_float(
            payload.get("holidays_carry_forward")
        )
    if "holiday_year_end" in payload:
        employee.holiday_year_end = _optional_date(
            payload.get("holiday_year_end"), field="holiday_year_end"
        )
    elif "annual_leave_restart" in payload:
        restart = _optional_date(
            payload.get("annual_leave_restart"), field="annual_leave_restart"
        )
        employee.holiday_year_end = (
            restart - dt.timedelta(days=1) if restart else None
        )
    if "holidays_remaining" in payload:
        employee.holidays_remaining = _optional_float(payload.get("holidays_remaining"))

    db.commit()
    db.refresh(employee)
    holiday_hours = _holiday_hours_by_employee(db, [employee.id]).get(employee.id, [])
    return _serialize_row(
        employee,
        entry,
        holiday_hours,
        include_rate=include_rate,
        as_of=start,
        period_start=start,
        period_end=end,
    )


def _active_staff_for_farm(db: Session, farm: str) -> list[Employee]:
    business = farm_business(farm)
    return sorted(
        db.scalars(
            select(Employee).where(
                Employee.status != EMPLOYEE_STATUS_ARCHIVED,
                Employee.business == business,
            )
        ).all(),
        key=_staff_sort_key,
    )


def _staff_sort_key(employee: Employee) -> tuple[int, int, str]:
    employment = employee.employment_type or EMPLOYMENT_TYPE_EMPLOYED
    employment_rank = 0 if employment == EMPLOYMENT_TYPE_EMPLOYED else 1
    pay_type = employee.pay_type or PAY_TYPE_HOURLY
    pay_rank = 0 if pay_type == PAY_TYPE_HOURLY else 1
    return (employment_rank, pay_rank, (employee.full_name or "").casefold())


def _holiday_hours_by_employee(
    db: Session, employee_ids: list[int]
) -> dict[int, list[EmployeeTimesheetEntry]]:
    if not employee_ids:
        return {}
    rows = db.scalars(
        select(EmployeeTimesheetEntry).where(
            EmployeeTimesheetEntry.employee_id.in_(employee_ids)
        )
    ).all()
    grouped: dict[int, list[EmployeeTimesheetEntry]] = {eid: [] for eid in employee_ids}
    for row in rows:
        grouped.setdefault(row.employee_id, []).append(row)
    return grouped


def _serialize_row(
    employee: Employee,
    entry: EmployeeTimesheetEntry | None,
    holiday_entries: list[EmployeeTimesheetEntry],
    *,
    include_rate: bool,
    as_of: dt.date | None = None,
    period_start: dt.date | None = None,
    period_end: dt.date | None = None,
) -> dict[str, Any]:
    as_of = as_of or dt.date.today()
    pay_type = employee.pay_type or PAY_TYPE_HOURLY
    hours_locked = pay_type == PAY_TYPE_SALARY
    rate_value = None
    rate_label = None
    salary_weekly = None
    if include_rate or hours_locked:
        amount, label = _display_rate(pay_type, decrypt_field(employee.pay_rate_enc))
        if hours_locked:
            salary_weekly = (
                round(amount / WEEKS_PER_YEAR, 2) if amount is not None else None
            )
        if include_rate:
            rate_value, rate_label = amount, label
    year_end = _employee_year_end(employee)
    holidays_taken = _holidays_taken_in_year(
        holiday_entries,
        year_end,
        as_of=as_of,
    )
    remaining, remaining_locked = _holidays_remaining(
        employee, holidays_taken, as_of
    )
    return {
        "employee_id": employee.id,
        "employee_number": employee.employee_number,
        "full_name": employee.full_name,
        "employment_type": employee.employment_type or EMPLOYMENT_TYPE_EMPLOYED,
        "employment_type_label": EMPLOYMENT_TYPE_LABELS.get(
            employee.employment_type or EMPLOYMENT_TYPE_EMPLOYED,
            "Employed",
        ),
        "pay_type": pay_type,
        "pay_type_label": PAY_TYPE_LABELS.get(pay_type, "Hourly"),
        "rate": rate_value,
        "rate_label": rate_label,
        "hours_locked": hours_locked,
        "hours_week_1": (
            salary_weekly
            if hours_locked
            else _num(entry.hours_week_1 if entry else None)
        ),
        "hours_week_2": (
            salary_weekly
            if hours_locked
            else _num(entry.hours_week_2 if entry else None)
        ),
        "holiday_hours": _num(entry.holiday_hours if entry else None),
        "dinner_break_hours": _num(entry.dinner_break_hours if entry else None),
        "mileage": _num(entry.mileage if entry else None),
        "bonus": _num(entry.bonus if entry else None),
        "loan_deduction": _num(entry.loan_deduction if entry else None),
        "other_deduction": _num(entry.other_deduction if entry else None),
        "accommodation_deduction": _accommodation_amount(
            employee, entry, period_start or as_of, period_end or as_of
        ),
        "other_remark": (entry.other_remark if entry else None) or "",
        "holidays_remaining": remaining,
        "remaining_locked": remaining_locked,
        "holidays_taken": holidays_taken,
        "holidays_carry_forward": _num(employee.holidays_carry_forward),
        "holiday_year_end": year_end.isoformat() if year_end else None,
        "annual_leave_days": (
            DEFAULT_ANNUAL_LEAVE_DAYS
            if employee.annual_leave_days is None
            else _num(employee.annual_leave_days)
        ),
    }


def leave_year_bounds(year_end: dt.date, as_of: dt.date) -> tuple[dt.date, dt.date]:
    """Holiday year containing as_of. year_end is the last day holiday can be taken."""
    this_end = _anniversary_on(as_of.year, year_end)
    end = this_end if as_of <= this_end else _anniversary_on(as_of.year + 1, year_end)
    start = _anniversary_on(end.year - 1, year_end) + dt.timedelta(days=1)
    return start, end


def _anniversary_on(year: int, template: dt.date) -> dt.date:
    try:
        return dt.date(year, template.month, template.day)
    except ValueError:
        last_day = calendar.monthrange(year, template.month)[1]
        return dt.date(year, template.month, last_day)


def _employee_year_end(employee: Employee) -> dt.date | None:
    if employee.holiday_year_end:
        return employee.holiday_year_end
    restart = getattr(employee, "annual_leave_restart", None)
    if restart:
        return restart - dt.timedelta(days=1)
    return None


def _accommodation_amount(
    employee: Employee,
    entry: EmployeeTimesheetEntry | None,
    period_start: dt.date | None,
    period_end: dt.date | None,
) -> float | None:
    if entry is not None and entry.accommodation_deduction is not None:
        return _num(entry.accommodation_deduction)
    amount = employee.accommodation_deduction
    if amount is None or period_start is None or period_end is None:
        return _num(amount)
    cadence = (employee.accommodation_cadence or ACCOMMODATION_CADENCE_WEEKLY).strip().lower()
    if cadence not in ACCOMMODATION_CADENCES:
        cadence = ACCOMMODATION_CADENCE_WEEKLY
    return _num(_accommodation_for_period(float(amount), cadence, period_start, period_end))


def _accommodation_for_period(
    amount: float,
    cadence: str,
    start: dt.date,
    end: dt.date,
) -> float:
    days = (end - start).days + 1
    if cadence == ACCOMMODATION_CADENCE_MONTHLY:
        month_days = calendar.monthrange(start.year, start.month)[1]
        if start.day == 1 and end.day == month_days and start.month == end.month:
            return round(amount, 2)
        return round(amount * days / month_days, 2)
    return round(amount * days / 7.0, 2)


def _holidays_taken_in_year(
    entries: list[EmployeeTimesheetEntry],
    year_end: dt.date | None,
    *,
    as_of: dt.date | None = None,
) -> float:
    year_start = None
    end = year_end
    as_of = as_of or dt.date.today()
    if year_end is not None:
        year_start, end = leave_year_bounds(year_end, as_of)
    total = 0.0
    for entry in entries:
        if year_start and (entry.period_start < year_start or entry.period_start > end):
            continue
        if entry.holiday_hours:
            total += float(entry.holiday_hours)
    return _num(total) or 0.0


def _holidays_remaining(
    employee: Employee,
    holidays_taken: float,
    as_of: dt.date,
) -> tuple[float | None, bool]:
    stored = _num(employee.holidays_remaining)
    year_end = _employee_year_end(employee)
    if stored is None and year_end is not None and as_of > year_end:
        entitlement = employee.annual_leave_days
        if entitlement is None:
            entitlement = DEFAULT_ANNUAL_LEAVE_DAYS
        return _num(float(entitlement) - float(holidays_taken or 0)), True
    return stored, False


def _display_rate(
    pay_type: str, raw: str | None
) -> tuple[float | None, str | None]:
    amount = _parse_money(raw)
    if amount is None:
        return None, None
    if pay_type == PAY_TYPE_SALARY:
        return round(amount, 2), f"£{amount:,.2f} / yr"
    return amount, f"£{amount:,.2f} / hr"


def _parse_money(raw: str | None) -> float | None:
    if not raw:
        return None
    cleaned = str(raw).replace("£", "").replace(",", "").strip()
    if not cleaned:
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def _optional_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise TimesheetError("Numeric fields must be numbers.") from exc
    if number < 0:
        raise TimesheetError("Numeric fields cannot be negative.")
    return number


def _num(value: float | None) -> float | None:
    if value is None:
        return None
    return float(value)


def _parse_date(value: dt.date | str | None, *, field: str) -> dt.date:
    if value is None or value == "":
        raise TimesheetError(f"{field} is required.")
    if isinstance(value, dt.date):
        return value
    try:
        return dt.date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise TimesheetError(f"{field} must be a valid date.") from exc


def _optional_date(value: dt.date | str | None, *, field: str) -> dt.date | None:
    if value in (None, ""):
        return None
    return _parse_date(value, field=field)


def _format_range(start: dt.date, end: dt.date) -> str:
    if start.month == end.month and start.year == end.year:
        return f"{start.day}-{end.day} {start.strftime('%b %Y')}"
    if start.year == end.year:
        return f"{start.day} {start.strftime('%b')} - {end.day} {end.strftime('%b %Y')}"
    return f"{start.day} {start.strftime('%b %Y')} - {end.day} {end.strftime('%b %Y')}"


def _uk_date(value: dt.date) -> str:
    return value.strftime("%d-%m-%Y")


def timesheet_xlsx_filename(start: dt.date, end: dt.date) -> str:
    return f"Time Sheet {_uk_date(start)} to {_uk_date(end)}.xlsx"


def _timesheet_xlsx_headers(period: dict[str, Any]) -> list[str]:
    hour_labels = [col["label"] for col in period.get("hours_columns") or []]
    return [
        "Emp. no.",
        "Name",
        "Employment type",
        "Pay type",
        "Rate",
        *hour_labels,
        "Holiday hours",
        "Dinner break (hrs)",
        "Mileage claimed",
        "Bonus",
        "Loan deduction",
        "Other deduction",
        "Accommodation deduction",
        "Other",
        "Holidays remaining",
        "Holidays taken",
        "Holidays carry forward",
        "Annual leave year end",
    ]


def _timesheet_xlsx_values(
    row: dict[str, Any], hours_columns: list[dict[str, Any]]
) -> list[Any]:
    year_end = row.get("holiday_year_end")
    if isinstance(year_end, str) and year_end:
        try:
            year_end = dt.date.fromisoformat(year_end)
        except ValueError:
            pass
    return [
        row.get("employee_number") or "",
        row.get("full_name") or "",
        row.get("employment_type_label") or "",
        row.get("pay_type_label") or "",
        row.get("rate_label") or "",
        *[row.get(col["key"]) for col in hours_columns],
        row.get("holiday_hours"),
        row.get("dinner_break_hours"),
        row.get("mileage"),
        row.get("bonus"),
        row.get("loan_deduction"),
        row.get("other_deduction"),
        row.get("accommodation_deduction"),
        row.get("other_remark") or "",
        row.get("holidays_remaining"),
        row.get("holidays_taken"),
        row.get("holidays_carry_forward"),
        year_end or "",
    ]


def _timesheet_row_fill(row: dict[str, Any]) -> PatternFill | None:
    employment = row.get("employment_type") or EMPLOYMENT_TYPE_EMPLOYED
    if employment == EMPLOYMENT_TYPE_SELF_EMPLOYED:
        return _SELF_EMPLOYED_FILL
    if (row.get("pay_type") or PAY_TYPE_HOURLY) == PAY_TYPE_SALARY:
        return _SALARY_FILL
    return None


def build_timesheet_xlsx(sheet: dict[str, Any]) -> bytes:
    period = sheet.get("period") or {}
    hours_columns = list(period.get("hours_columns") or [])
    headers = _timesheet_xlsx_headers(period)
    last_col = get_column_letter(len(headers))
    money_headers = {
        "Mileage claimed",
        "Bonus",
        "Loan deduction",
        "Other deduction",
        "Accommodation deduction",
    }
    hours_headers = {
        "Holiday hours",
        "Dinner break (hrs)",
        "Holidays remaining",
        "Holidays taken",
        "Holidays carry forward",
        *[col["label"] for col in hours_columns],
    }
    hour_labels = {col["label"] for col in hours_columns}

    wb = Workbook()
    ws = wb.active
    ws.title = "Time Sheet"

    start = dt.date.fromisoformat(period["start"]) if period.get("start") else None
    end = dt.date.fromisoformat(period["end"]) if period.get("end") else None
    farm_name = sheet.get("farm_label") or sheet.get("business") or ""
    range_label = ""
    if start and end:
        range_label = f"{_uk_date(start)} to {_uk_date(end)}"
    subtitle = " · ".join(
        part for part in (farm_name, period.get("label") or "", range_label) if part
    )

    ws.merge_cells(f"A1:{last_col}1")
    ws.merge_cells(f"A2:{last_col}2")
    title_cell = ws["A1"]
    title_cell.value = "Time Sheet"
    title_cell.font = _TITLE_FONT
    title_cell.fill = _TITLE_FILL
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    subtitle_cell = ws["A2"]
    subtitle_cell.value = subtitle
    subtitle_cell.font = _SUBTITLE_FONT
    subtitle_cell.fill = _SUBTITLE_FILL
    subtitle_cell.alignment = Alignment(horizontal="left", vertical="center")
    for cell in ws[1] + ws[2]:
        cell.border = _THIN_BORDER
        if cell.coordinate == "A1":
            cell.fill = _TITLE_FILL
        elif cell.coordinate == "A2":
            cell.fill = _SUBTITLE_FILL
        else:
            cell.fill = _TITLE_FILL if cell.row == 1 else _SUBTITLE_FILL

    ws.append(headers)
    header_row = ws[3]
    for cell in header_row:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = _THIN_BORDER
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 32

    for row in sheet.get("rows") or []:
        values = _timesheet_xlsx_values(row, hours_columns)
        ws.append(values)
        excel_row = ws.max_row
        fill = _timesheet_row_fill(row)
        is_salary = (row.get("pay_type") or PAY_TYPE_HOURLY) == PAY_TYPE_SALARY
        for index, cell in enumerate(ws[excel_row], start=1):
            header = headers[index - 1]
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(vertical="center")
            if fill is not None:
                cell.fill = fill
            if header in money_headers and isinstance(cell.value, (int, float)):
                cell.number_format = _MONEY_FORMAT
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif (
                is_salary
                and header in hour_labels
                and isinstance(cell.value, (int, float))
            ):
                cell.number_format = _MONEY_FORMAT
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif header in hours_headers and isinstance(cell.value, (int, float)):
                cell.number_format = _HOURS_FORMAT
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif header == "Annual leave year end" and isinstance(cell.value, dt.date):
                cell.number_format = _DATE_FORMAT
                cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [12, 24, 16, 12, 14]
    widths.extend([14] * len(hours_columns))
    widths.extend([13, 14, 14, 12, 14, 14, 18, 22, 14, 13, 16, 16])
    for index, width in enumerate(widths[: len(headers)], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:{last_col}{max(ws.max_row, 3)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:3"
    ws.page_setup.horizontalCentered = True

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

