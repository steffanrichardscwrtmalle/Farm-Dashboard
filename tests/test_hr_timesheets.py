"""Time sheet pay periods, active staff listing, and row saves."""

from __future__ import annotations

import datetime as dt

from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.models import (
    EMPLOYEE_STATUS_ACTIVE,
    EMPLOYEE_STATUS_ARCHIVED,
    EMPLOYEE_STATUS_ONBOARDING,
    EMPLOYEE_STATUS_PENDING_SIGNATURE,
    EMPLOYMENT_TYPE_EMPLOYED,
    EMPLOYMENT_TYPE_SELF_EMPLOYED,
    PAY_TYPE_HOURLY,
    PAY_TYPE_SALARY,
    Base,
    Employee,
)
from app.services.crypto_fields import encrypt_field
from app.services.timesheet_service import (
    TimesheetError,
    build_timesheet_xlsx,
    current_holiday_year_end,
    current_period_start,
    leave_year_bounds,
    list_periods,
    list_timesheet,
    period_containing,
    save_timesheet_row,
    timesheet_xlsx_filename,
)


@pytest.fixture()
def db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine, autoflush=False, autocommit=False)()
    try:
        yield session
    finally:
        session.close()


def _employee(db, **overrides):
    data = {
        "business": "Cwrt Malle Ltd",
        "employment_type": EMPLOYMENT_TYPE_EMPLOYED,
        "employee_number": "CM001",
        "full_name": "Alex Farmhand",
        "email": "alex@test.local",
        "pay_type": PAY_TYPE_HOURLY,
        "pay_rate_enc": encrypt_field("12.50"),
        "role_title": "Farm Worker",
        "start_date": dt.date(2026, 1, 1),
        "status": EMPLOYEE_STATUS_ACTIVE,
    }
    data.update(overrides)
    row = Employee(**data)
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def test_cm_fortnight_starts_7_sep_2026():
    start, end = period_containing("CM", dt.date(2026, 9, 7))
    assert start == dt.date(2026, 9, 7)
    assert end == dt.date(2026, 9, 20)
    start, end = period_containing("CM", dt.date(2026, 9, 21))
    assert start == dt.date(2026, 9, 21)
    assert end == dt.date(2026, 10, 4)


def test_gad_month_starts_1_sep_2026():
    start, end = period_containing("GAD", dt.date(2026, 9, 21))
    assert start == dt.date(2026, 9, 1)
    assert end == dt.date(2026, 9, 30)
    start, end = period_containing("GAD", dt.date(2026, 10, 8))
    assert start == dt.date(2026, 10, 1)
    assert end == dt.date(2026, 10, 31)


def test_period_lists_include_cycle_starts():
    as_of = dt.date(2026, 9, 21)
    cm = list_periods("CM", as_of=as_of)
    gad = list_periods("GAD", as_of=as_of)
    assert cm[0]["start"] == "2026-09-07"
    assert cm[0]["end"] == "2026-09-20"
    assert [col["key"] for col in cm[0]["hours_columns"]] == [
        "hours_week_1",
        "hours_week_2",
    ]
    assert gad[0]["start"] == "2026-09-01"
    assert gad[0]["label"] == "September 2026"
    assert [col["key"] for col in gad[0]["hours_columns"]] == ["hours_week_1"]
    assert current_period_start("CM", as_of=as_of) == dt.date(2026, 9, 21)
    assert current_period_start("GAD", as_of=as_of) == dt.date(2026, 9, 1)


def test_lists_current_staff_including_onboarding(db):
    _employee(db, employee_number="CM001", full_name="Active Employed")
    _employee(
        db,
        employee_number="CM002",
        full_name="Active Contractor",
        email="sam@test.local",
        employment_type=EMPLOYMENT_TYPE_SELF_EMPLOYED,
    )
    _employee(
        db,
        employee_number="CM003",
        full_name="Onboarding",
        email="on@test.local",
        status=EMPLOYEE_STATUS_ONBOARDING,
    )
    _employee(
        db,
        employee_number="CM005",
        full_name="Pending Signature",
        email="pending@test.local",
        status=EMPLOYEE_STATUS_PENDING_SIGNATURE,
    )
    _employee(
        db,
        employee_number="CM004",
        full_name="Archived",
        email="left@test.local",
        status=EMPLOYEE_STATUS_ARCHIVED,
    )
    _employee(
        db,
        employee_number="GAD001",
        full_name="GAD Worker",
        email="gad@test.local",
        business="Green Acre Dairy Ltd",
    )
    sheet = list_timesheet(
        db, "CM", period_start=dt.date(2026, 9, 7), include_rate=True
    )
    names = [row["full_name"] for row in sheet["rows"]]
    assert names == [
        "Active Employed",
        "Onboarding",
        "Pending Signature",
        "Active Contractor",
    ]
    assert [row["employment_type_label"] for row in sheet["rows"]] == [
        "Employed",
        "Employed",
        "Employed",
        "Self-employed",
    ]
    assert sheet["rows"][0]["rate_label"] == "£12.50 / hr"
    assert sheet["rows"][0]["holiday_days"] == 0
    assert sheet["rows"][0]["holiday_hours"] == 0
    assert sheet["rows"][0]["holiday_hours_per_day"] == 8
    assert sheet["cadence"] == "fortnightly"


def test_sorts_employed_hourly_then_salary_then_self_employed(db):
    _employee(
        db,
        employee_number="CM020",
        full_name="Zoe Hourly",
        email="zoe@test.local",
        pay_type=PAY_TYPE_HOURLY,
    )
    _employee(
        db,
        employee_number="CM021",
        full_name="Amy Salary",
        email="amy@test.local",
        pay_type=PAY_TYPE_SALARY,
    )
    _employee(
        db,
        employee_number="CM022",
        full_name="Ben Hourly",
        email="ben@test.local",
        pay_type=PAY_TYPE_HOURLY,
    )
    _employee(
        db,
        employee_number="CM023",
        full_name="Ada Contractor",
        email="ada@test.local",
        employment_type=EMPLOYMENT_TYPE_SELF_EMPLOYED,
        pay_type=PAY_TYPE_HOURLY,
    )
    _employee(
        db,
        employee_number="CM024",
        full_name="Cal Contractor",
        email="cal@test.local",
        employment_type=EMPLOYMENT_TYPE_SELF_EMPLOYED,
        pay_type=PAY_TYPE_SALARY,
    )
    sheet = list_timesheet(db, "CM", period_start=dt.date(2026, 9, 7))
    assert [row["full_name"] for row in sheet["rows"]] == [
        "Ben Hourly",
        "Zoe Hourly",
        "Amy Salary",
        "Ada Contractor",
        "Cal Contractor",
    ]


def test_salary_rate_shown_annually(db):
    _employee(
        db,
        pay_type=PAY_TYPE_SALARY,
        pay_rate_enc=encrypt_field("26000"),
    )
    sheet = list_timesheet(
        db, "CM", period_start=dt.date(2026, 9, 7), include_rate=True
    )
    assert sheet["rows"][0]["pay_type_label"] == "Salary"
    assert sheet["rows"][0]["rate"] == 26000.0
    assert sheet["rows"][0]["rate_label"] == "£26,000.00 / yr"
    assert sheet["rows"][0]["hours_locked"] is True
    assert sheet["rows"][0]["hours_week_1"] == 500.0
    assert sheet["rows"][0]["hours_week_2"] == 500.0


def test_salary_hours_cannot_be_overwritten(db):
    staff = _employee(
        db,
        pay_type=PAY_TYPE_SALARY,
        pay_rate_enc=encrypt_field("26000"),
    )
    saved = save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": staff.id,
            "period_start": dt.date(2026, 9, 7),
            "hours_week_1": 40,
            "hours_week_2": 38,
            "holiday_days": 1,
        },
        include_rate=True,
    )
    assert saved["hours_locked"] is True
    assert saved["hours_week_1"] == 500.0
    assert saved["hours_week_2"] == 500.0
    assert saved["holiday_days"] == 1
    assert saved["holiday_hours"] == 8
    sheet = list_timesheet(db, "CM", period_start=dt.date(2026, 9, 7))
    assert sheet["rows"][0]["hours_week_1"] == 500.0
    assert sheet["rows"][0]["hours_week_2"] == 500.0


def test_save_hours_and_holiday_year_total(db):
    staff = _employee(db, holiday_year_end=dt.date(2027, 3, 31))
    saved = save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": staff.id,
            "period_start": dt.date(2026, 9, 7),
            "hours_week_1": 40,
            "hours_week_2": 32,
            "holiday_days": 1,
            "dinner_break_hours": 5,
            "mileage": 22.5,
            "bonus": 50,
            "loan_deduction": 10,
            "other_deduction": 5,
            "accommodation_deduction": 80,
            "other_remark": "Covered weekend milking",
            "holidays_remaining": 72,
            "holidays_carry_forward": 8,
            "holiday_year_end": dt.date(2028, 1, 1),
        },
        include_rate=True,
    )
    assert saved["hours_week_1"] == 40
    assert saved["hours_week_2"] == 32
    assert saved["holiday_days"] == 1
    assert saved["holiday_hours"] == 8
    assert saved["holidays_taken"] == 1
    assert saved["holidays_remaining"] == 72
    assert saved["other_remark"] == "Covered weekend milking"
    assert saved["holiday_year_end"] == "2027-03-31"

    save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": staff.id,
            "period_start": dt.date(2026, 9, 21),
            "holiday_days": 2,
            "holiday_year_end": dt.date(2027, 3, 31),
            "holidays_remaining": 56,
            "holidays_carry_forward": 8,
        },
    )
    sheet = list_timesheet(
        db, "CM", period_start=dt.date(2026, 9, 21), include_rate=False
    )
    assert sheet["rows"][0]["holidays_taken"] == 3
    assert sheet["rows"][0]["holiday_days"] == 2
    assert sheet["rows"][0]["holiday_hours"] == 16
    assert sheet["rows"][0]["rate"] is None


def test_rejects_wrong_farm_and_invalid_period(db):
    staff = _employee(db, business="Green Acre Dairy Ltd")
    with pytest.raises(TimesheetError, match="does not belong"):
        save_timesheet_row(
            db,
            "CM",
            {"employee_id": staff.id, "period_start": dt.date(2026, 9, 7)},
        )
    with pytest.raises(TimesheetError, match="first day of a pay period"):
        save_timesheet_row(
            db,
            "GAD",
            {"employee_id": staff.id, "period_start": dt.date(2026, 9, 7)},
        )


def test_can_save_onboarding_staff_but_not_archived(db):
    onboarding = _employee(
        db,
        employee_number="CM010",
        email="new@test.local",
        status=EMPLOYEE_STATUS_ONBOARDING,
    )
    saved = save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": onboarding.id,
            "period_start": dt.date(2026, 9, 7),
            "hours_week_1": 20,
        },
    )
    assert saved["hours_week_1"] == 20

    archived = _employee(
        db,
        employee_number="CM011",
        email="gone@test.local",
        status=EMPLOYEE_STATUS_ARCHIVED,
    )
    with pytest.raises(TimesheetError, match="Archived"):
        save_timesheet_row(
            db,
            "CM",
            {"employee_id": archived.id, "period_start": dt.date(2026, 9, 7)},
        )


def test_leave_year_bounds_from_september_year_end():
    start, end = leave_year_bounds(dt.date(2026, 9, 30), dt.date(2026, 9, 21))
    assert start == dt.date(2025, 10, 1)
    assert end == dt.date(2026, 9, 30)
    start, end = leave_year_bounds(dt.date(2026, 9, 30), dt.date(2026, 9, 30))
    assert start == dt.date(2025, 10, 1)
    assert end == dt.date(2026, 9, 30)
    start, end = leave_year_bounds(dt.date(2026, 9, 30), dt.date(2026, 10, 1))
    assert start == dt.date(2026, 10, 1)
    assert end == dt.date(2027, 9, 30)


def test_holiday_year_end_moves_forward_one_year_after_it_passes():
    end = dt.date(2026, 8, 24)
    assert current_holiday_year_end(end, dt.date(2026, 8, 24)) == dt.date(2026, 8, 24)
    assert current_holiday_year_end(end, dt.date(2026, 8, 25)) == dt.date(2027, 8, 24)
    assert current_holiday_year_end(end, dt.date(2028, 9, 1)) == dt.date(2029, 8, 24)


def test_timesheet_year_end_comes_from_directory_and_rolls(db):
    staff = _employee(db, holiday_year_end=dt.date(2026, 9, 20))
    before = list_timesheet(db, "CM", period_start=dt.date(2026, 9, 7))
    assert before["rows"][0]["holiday_year_end"] == "2026-09-20"
    after = list_timesheet(db, "CM", period_start=dt.date(2026, 9, 21))
    assert after["rows"][0]["holiday_year_end"] == "2027-09-20"
    db.refresh(staff)
    assert staff.holiday_year_end == dt.date(2026, 9, 20)


def test_remaining_resets_the_day_after_year_end(db):
    staff = _employee(
        db,
        holiday_year_end=dt.date(2026, 9, 30),
        annual_leave_days=28,
        holidays_remaining=None,
    )
    before = save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": staff.id,
            "period_start": dt.date(2026, 9, 21),
            "holiday_days": 4,
        },
    )
    assert before["remaining_locked"] is False
    assert before["holidays_remaining"] is None
    assert before["holiday_year_end"] == "2026-09-30"

    after = save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": staff.id,
            "period_start": dt.date(2026, 10, 5),
            "holiday_days": 8,
        },
    )
    assert after["remaining_locked"] is True
    assert after["holidays_taken"] == 8
    assert after["holidays_remaining"] == 20
    assert after["holiday_year_end"] == "2027-09-30"


def test_imported_remaining_is_kept_after_year_end(db):
    staff = _employee(
        db,
        holiday_year_end=dt.date(2026, 8, 24),
        annual_leave_days=28,
        holidays_remaining=19,
    )
    saved = save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": staff.id,
            "period_start": dt.date(2026, 9, 7),
        },
    )
    assert saved["remaining_locked"] is False
    assert saved["holidays_remaining"] == 19


def test_timesheet_uses_staff_accommodation_default(db):
    staff = _employee(db, accommodation_deduction=80, accommodation_cadence="weekly")
    sheet = list_timesheet(db, "CM", period_start=dt.date(2026, 9, 7))
    assert sheet["rows"][0]["employee_id"] == staff.id
    assert sheet["rows"][0]["accommodation_deduction"] == 160


def test_monthly_accommodation_on_fortnight_is_prorated(db):
    staff = _employee(
        db,
        accommodation_deduction=80,
        accommodation_cadence="monthly",
    )
    sheet = list_timesheet(db, "CM", period_start=dt.date(2026, 9, 7))
    assert sheet["rows"][0]["accommodation_deduction"] == 37.33


def test_monthly_accommodation_on_gad_month_is_full_amount(db):
    staff = _employee(
        db,
        business="Green Acre Dairy Ltd",
        email="gad-acc@test.local",
        employee_number="GAD001",
        accommodation_deduction=80,
        accommodation_cadence="monthly",
    )
    sheet = list_timesheet(db, "GAD", period_start=dt.date(2026, 9, 1))
    assert sheet["rows"][0]["employee_id"] == staff.id
    assert sheet["rows"][0]["accommodation_deduction"] == 80


def test_timesheet_xlsx_filename_uses_uk_date_range():
    assert (
        timesheet_xlsx_filename(dt.date(2026, 9, 7), dt.date(2026, 9, 20))
        == "Time Sheet 07-09-2026 to 20-09-2026.xlsx"
    )
    assert (
        timesheet_xlsx_filename(dt.date(2026, 9, 1), dt.date(2026, 9, 30))
        == "Time Sheet 01-09-2026 to 30-09-2026.xlsx"
    )


def test_timesheet_xlsx_is_formatted_workbook(db):
    hourly = _employee(db, employee_number="CM010", full_name="Alex Farmhand")
    salary = _employee(
        db,
        employee_number="CM020",
        full_name="Sam Salary",
        email="sam@test.local",
        pay_type=PAY_TYPE_SALARY,
        pay_rate_enc=encrypt_field("28000"),
    )
    contractor = _employee(
        db,
        employee_number="CM030",
        full_name="Pat Contractor",
        email="pat@test.local",
        employment_type=EMPLOYMENT_TYPE_SELF_EMPLOYED,
    )
    save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": hourly.id,
            "period_start": dt.date(2026, 9, 7),
            "hours_week_1": 40,
            "hours_week_2": 38,
            "holiday_days": 1.25,
            "mileage": 12.5,
            "other_remark": "Covered relief",
        },
    )
    sheet = list_timesheet(
        db, "CM", period_start=dt.date(2026, 9, 7), include_rate=True
    )
    content = build_timesheet_xlsx(sheet)
    wb = load_workbook(BytesIO(content))
    ws = wb.active
    assert ws.title == "Time Sheet"
    assert ws["A1"].value == "Time Sheet"
    assert "07-09-2026 to 20-09-2026" in str(ws["A2"].value)
    assert "Cwrt Malle" in str(ws["A2"].value)
    headers = [cell.value for cell in ws[3]]
    assert headers[:5] == [
        "Emp. no.",
        "Name",
        "Employment type",
        "Pay type",
        "Rate",
    ]
    assert "Hours 7-13 Sep 2026" in headers
    assert "Hours 14-20 Sep 2026" in headers
    days_col = headers.index("Days holiday")
    hours_col_h = headers.index("Holiday hours")
    assert days_col == hours_col_h - 1
    assert ws["A1"].font.bold is True
    assert ws["A3"].font.bold is True
    assert ws.freeze_panes == "C4"
    names = [ws.cell(row=index, column=2).value for index in range(4, ws.max_row + 1)]
    assert names == ["Alex Farmhand", "Sam Salary", "Pat Contractor"]
    hours_col = headers.index("Hours 7-13 Sep 2026") + 1
    assert ws.cell(row=4, column=hours_col).value == 40
    assert ws.cell(row=4, column=hours_col).number_format == "0.00"
    assert ws.cell(row=5, column=hours_col).value == 538.46
    assert ws.cell(row=5, column=hours_col).number_format == "£#,##0.00"
    days_holiday_col = headers.index("Days holiday") + 1
    holiday_hours_col = headers.index("Holiday hours") + 1
    assert days_holiday_col == holiday_hours_col - 1
    assert ws.cell(row=4, column=days_holiday_col).value == 1.25
    assert ws.cell(row=4, column=holiday_hours_col).value == 10
    assert ws.cell(row=4, column=days_holiday_col).number_format == "0.00"
    assert ws.cell(row=4, column=holiday_hours_col).number_format == "0.00"
    mileage_col = headers.index("Mileage claimed") + 1
    assert ws.cell(row=4, column=mileage_col).value == 12.5
    assert ws.cell(row=4, column=mileage_col).number_format == "£#,##0.00"
    salary_fill = ws.cell(row=5, column=2).fill.fgColor.rgb
    contractor_fill = ws.cell(row=6, column=2).fill.fgColor.rgb
    assert salary_fill.endswith("D4EDDA")
    assert contractor_fill.endswith("EFE6F7")
    assert hourly.id and salary.id and contractor.id


def test_holiday_hours_equal_days_times_hours_per_day(db):
    staff = _employee(db)
    blank = list_timesheet(db, "CM", period_start=dt.date(2026, 9, 7))
    assert blank["rows"][0]["holiday_days"] == 0
    assert blank["rows"][0]["holiday_hours"] == 0
    assert blank["rows"][0]["holiday_hours_per_day"] == 8

    one_day = save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": staff.id,
            "period_start": dt.date(2026, 9, 7),
            "holiday_days": 1,
        },
    )
    assert one_day["holiday_days"] == 1
    assert one_day["holiday_hours"] == 8
    assert one_day["holidays_taken"] == 1

    quarter = save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": staff.id,
            "period_start": dt.date(2026, 9, 7),
            "holiday_days": 0.25,
        },
    )
    assert quarter["holiday_days"] == 0.25
    assert quarter["holiday_hours"] == 2

    rounded = save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": staff.id,
            "period_start": dt.date(2026, 9, 7),
            "holiday_days": 1.3,
        },
    )
    assert rounded["holiday_days"] == 1.25
    assert rounded["holiday_hours"] == 10


def test_custom_holiday_hours_per_day_is_used_on_timesheet(db):
    staff = _employee(db, holiday_hours_per_day=7.5)
    saved = save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": staff.id,
            "period_start": dt.date(2026, 9, 7),
            "holiday_days": 1,
        },
    )
    assert saved["holiday_hours_per_day"] == 7.5
    assert saved["holiday_days"] == 1
    assert saved["holiday_hours"] == 7.5


def test_hours_per_day_holiday_is_editable_for_employed_staff(db):
    from app.services.hr_service import (
        HRServiceError,
        get_staff_detail,
        update_employee_holiday_hours_per_day,
    )

    staff = _employee(db)
    detail = get_staff_detail(db, staff.id)
    assert detail["holiday_hours_per_day"] == 8
    updated = update_employee_holiday_hours_per_day(db, staff.id, 6)
    assert updated["holiday_hours_per_day"] == 6
    db.refresh(staff)
    assert staff.holiday_hours_per_day == 6
    saved = save_timesheet_row(
        db,
        "CM",
        {
            "employee_id": staff.id,
            "period_start": dt.date(2026, 9, 7),
            "holiday_days": 0.5,
        },
    )
    assert saved["holiday_hours"] == 3

    contractor = _employee(
        db,
        employee_number="CM099",
        email="contractor@test.local",
        employment_type=EMPLOYMENT_TYPE_SELF_EMPLOYED,
    )
    with pytest.raises(HRServiceError, match="employed staff"):
        update_employee_holiday_hours_per_day(db, contractor.id, 8)

