"""Haulier XLSX date-header parsing (typos and truncated years)."""

from __future__ import annotations

import datetime as dt
import io
from pathlib import Path

import openpyxl

from app.services.haulier_xlsx import _parse_date_header, parse_haulier_xlsx


def test_two_digit_year_uses_previous_header_century() -> None:
    prev = dt.date(2026, 7, 24)
    assert _parse_date_header("Saturday 25th July 26", prev) == dt.date(2026, 7, 25)
    assert _parse_date_header("Sunday 26th July 26", prev) == dt.date(2026, 7, 26)


def test_month_typo_jult_still_parses() -> None:
    prev = dt.date(2026, 7, 29)
    assert _parse_date_header("Thursday 30th Jult 2026 ", prev) == dt.date(2026, 7, 30)


def test_weekday_typo_firday_is_the_second() -> None:
    prev = dt.date(2026, 10, 1)
    assert _parse_date_header("Firday 2nd October 2026", prev) == dt.date(2026, 10, 2)


def test_firday_header_moves_following_loads_off_the_first() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active

    def put(row: int, customer: str, volume: int | None = None, sample: str | None = None) -> None:
        ws.cell(row, 1, customer)
        if volume is not None:
            ws.cell(row, 14, volume)
            ws.cell(row, 17, dt.time(4, 30))
        if sample is not None:
            ws.cell(row, 23, sample)

    put(1, "Thursday 1st October 2026")
    put(2, "Cwrt Malle", 1000, "085")
    put(3, "Firday 2nd October 2026")
    put(4, "Cwrt Malle", 2000, "089")
    buf = io.BytesIO()
    wb.save(buf)

    by_day: dict[dt.date, list[str]] = {}
    for row in parse_haulier_xlsx(buf.getvalue())["rows"]:
        by_day.setdefault(row["collection_date"], []).append(row["sample_id"])

    assert by_day[dt.date(2026, 10, 1)] == ["085"]
    assert by_day[dt.date(2026, 10, 2)] == ["089"]


def test_july_sheet_splits_24_through_27() -> None:
    path = Path(__file__).resolve().parents[1] / "Cwrt Malle - 07 July 2026.xlsx"
    if not path.is_file():
        return
    rows = parse_haulier_xlsx(path.read_bytes())["rows"]
    by_day: dict[dt.date, list] = {}
    for row in rows:
        by_day.setdefault(row["collection_date"], []).append(row)

    assert len(by_day[dt.date(2026, 7, 24)]) == 3
    assert {r["sample_id"] for r in by_day[dt.date(2026, 7, 24)]} == {"090", "091", "092"}

    assert len(by_day[dt.date(2026, 7, 25)]) == 4
    assert {r["sample_id"] for r in by_day[dt.date(2026, 7, 25)]} == {
        "093",
        "094",
        "095",
        "096",
    }

    assert len(by_day[dt.date(2026, 7, 26)]) == 3
    assert {r["sample_id"] for r in by_day[dt.date(2026, 7, 26)]} == {
        "097",
        "098",
        "099",
    }

    assert len(by_day[dt.date(2026, 7, 27)]) == 4
    assert sum(r["volume_litres"] or 0 for r in by_day[dt.date(2026, 7, 24)]) == 84402
    assert sum(r["volume_litres"] or 0 for r in by_day[dt.date(2026, 7, 25)]) == 112768
    assert sum(r["volume_litres"] or 0 for r in by_day[dt.date(2026, 7, 26)]) == 84043
