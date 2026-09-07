"""Tests for Feedlync Loaded Mixes feed usage."""

from __future__ import annotations

import datetime as dt
from io import BytesIO
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker

from app.models import Base, FeedUsageRecord
from app.services.feed_usage import (
    build_usage_report,
    build_usage_xlsx,
    get_usage_report,
    import_feed_usage,
    resolve_usage_month,
)
from app.services.feed_usage_settings import (
    assigned_ration_names,
    ingredient_inclusion_lookup,
    ration_farm_lookup,
    save_ingredient_assignment,
    save_ration_assignment,
    seed_ration_assignments_if_empty,
)
from app.services.feedlync_api import (
    _LOADED_MIX_INGREDIENT_PATH,
    _parse_ingredient_spends,
    _parse_ingredient_spends_by_farm,
    _utc_month_range,
    farm_for_usage_ration,
    fetch_loaded_mix_ingredient_usage,
    month_bounds,
    previous_calendar_month,
)


AUG_START = dt.date(2026, 8, 1)
AUG_END = dt.date(2026, 8, 31)


@pytest.fixture
def db() -> Session:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine, autoflush=False, autocommit=False)()
    try:
        yield session
    finally:
        session.close()


def test_month_bounds_full_calendar_month() -> None:
    start, end = month_bounds(2026, 8)
    assert start == AUG_START
    assert end == AUG_END


def test_previous_calendar_month_from_september() -> None:
    assert previous_calendar_month(dt.date(2026, 9, 7)) == AUG_START


def test_resolve_usage_month_defaults_to_last_month() -> None:
    assert resolve_usage_month(None, today=dt.date(2026, 9, 7)) == AUG_START
    assert resolve_usage_month("2026-08", today=dt.date(2026, 9, 7)) == AUG_START


def test_utc_month_range_matches_feedlync_last_month_encoding() -> None:
    start, end = _utc_month_range(AUG_START, AUG_END)
    assert start == "2026-08-01T00:00:00.000Z"
    assert end == "2026-09-01T00:00:00.000Z"


def test_parse_ingredient_spends_sorts_and_merges() -> None:
    rows = _parse_ingredient_spends(
        [
            {
                "ingredientName": "Rape Meal",
                "quantity": 200000,
                "drymatterQuantity": 180000,
                "cost": 50000,
            },
            {
                "ingredientName": "Rape Meal",
                "quantity": 322593,
                "drymatterQuantity": 290000,
                "cost": 70000.25,
            },
            {
                "ingredientName": "Grass Silage",
                "quantity": 100,
                "drymatterQuantity": 30,
                "cost": 10,
            },
        ]
    )
    assert rows[0]["ingredient_name"] == "Rape Meal"
    assert rows[0]["as_fed_kg"] == 522593
    assert rows[0]["dm_kg"] == 470000
    assert rows[0]["cost"] == 120000.25
    assert rows[1]["ingredient_name"] == "Grass Silage"


def test_build_usage_report_totals_and_daily_averages() -> None:
    report = build_usage_report(
        [
            {
                "ingredient_name": "Grass Silage",
                    "as_fed_kg": 6_398_038,
                "dm_kg": 3_404_860,
                "cost": 872_640.75,
            },
            {
                "ingredient_name": "Rape Meal",
                "as_fed_kg": 522_593,
                "dm_kg": 470_000,
                "cost": 150_000,
            },
        ],
        period_start=AUG_START,
        period_end=AUG_END,
        farm="GAD",
    )
    assert report["farm"] == "GAD"
    assert report["farm_label"] == "Green Acre Dairy"
    assert "Coomb Milker Premix" in report["rations"]
    assert report["ingredients"][0]["ingredient_name"] == "Grass Silage"
    assert report["ingredients"][1]["ingredient_name"] == "Rape Meal"
    assert report["ingredients"][1]["as_fed_kg"] == 522_593
    assert report["totals"]["as_fed_kg"] == 6_920_631
    assert report["totals"]["dm_kg"] == 3_874_860
    assert report["totals"]["cost"] == 1_022_640.75
    assert report["days"] == 31
    assert report["averages"]["as_fed_kg"] == round(6_920_631 / 31, 2)
    assert report["source"] == "Loaded Mixes → By Ingredient"


def test_import_replaces_only_selected_month(db: Session) -> None:
    july_start, july_end = month_bounds(2026, 7)
    db.add(
        FeedUsageRecord(
            period_start=july_start,
            period_end=july_end,
            farm="CM",
            ingredient_name="Keep Me",
            as_fed_kg=1,
            dm_kg=1,
            cost=1,
        )
    )
    db.commit()

    payload_rows = [
        {
            "farm": "GAD",
            "ingredient_name": "Rape Meal",
            "as_fed_kg": 522593,
            "dm_kg": 470000,
            "cost": 150000,
        },
        {
            "farm": "CM",
            "ingredient_name": "Rape Meal",
            "as_fed_kg": 1000,
            "dm_kg": 900,
            "cost": 250,
        },
    ]
    with patch(
        "app.services.feed_usage.fetch_loaded_mix_ingredient_usage",
        return_value=payload_rows,
    ):
        import_feed_usage(db, month=AUG_START)

    names = {
        (row.period_start, row.farm, row.ingredient_name)
        for row in db.scalars(select(FeedUsageRecord)).all()
    }
    assert (july_start, "CM", "Keep Me") in names
    assert (AUG_START, "GAD", "Rape Meal") in names
    assert (AUG_START, "CM", "Rape Meal") in names
    august = get_usage_report(db, month=AUG_START, farm="GAD")
    assert august["ingredients"][0]["as_fed_kg"] == 522593
    assert august["row_count"] == 1
    assert august["farm"] == "GAD"
    cm = get_usage_report(db, month=AUG_START, farm="CM")
    assert cm["ingredients"][0]["as_fed_kg"] == 1000
    assert cm["row_count"] == 1


def test_fetch_uses_loaded_mixes_endpoint_not_fed_mixes() -> None:
    response = MagicMock()
    response.status_code = 200
    response.json.return_value = [
        {
            "ingredientName": "Rape Meal",
            "quantity": 999999,
            "drymatterQuantity": 999999,
            "cost": 999999,
            "rations": [
                {
                    "rationName": "Coomb Milker Premix",
                    "quantity": 522593,
                    "drymatterQuantity": 470000,
                    "cost": 150000,
                }
            ],
        }
    ]
    response.raise_for_status.return_value = None
    client = MagicMock()
    client.get.return_value = response
    client.__enter__.return_value = client
    client.__exit__.return_value = False

    with (
        patch("app.services.feedlync_api.httpx.Client", return_value=client),
        patch(
            "app.services.feedlync_api._authenticate_farms",
            return_value=("token", ["farm-1"]),
        ),
        patch(
            "app.services.feed_usage_settings.ration_farm_lookup",
            return_value={"coomb milker premix": "GAD"},
        ),
        patch(
            "app.services.feed_usage_settings.ingredient_inclusion_lookup",
            return_value=None,
        ),
    ):
        rows = fetch_loaded_mix_ingredient_usage(
            MagicMock(), period_start=AUG_START, period_end=AUG_END
        )

    assert rows[0]["ingredient_name"] == "Rape Meal"
    assert rows[0]["farm"] == "GAD"
    assert rows[0]["as_fed_kg"] == 522593
    url = client.get.call_args.args[0]
    assert _LOADED_MIX_INGREDIENT_PATH in url
    assert "fedmix" not in url.lower()
    params = client.get.call_args.kwargs["params"]
    assert params["from"] == "2026-08-01T00:00:00.000Z"
    assert params["to"] == "2026-09-01T00:00:00.000Z"


def test_farm_for_usage_ration_matches_named_rations() -> None:
    assert farm_for_usage_ration("Coomb Milker Premix") == "GAD"
    assert farm_for_usage_ration("coomb far off") == "GAD"
    assert farm_for_usage_ration("Cwrt Milkers") == "CM"
    assert farm_for_usage_ration("Coomb Milkers") is None
    assert farm_for_usage_ration("Cwrt Pre-Bullers") is None


def test_parse_ingredient_spends_by_farm_ignores_parent_totals() -> None:
    rows = _parse_ingredient_spends_by_farm(
        [
            {
                "ingredientName": "Rape Meal",
                "quantity": 999999,
                "drymatterQuantity": 999999,
                "cost": 999999,
                "rations": [
                    {
                        "rationName": "Coomb Milker Premix",
                        "quantity": 100,
                        "drymatterQuantity": 90,
                        "cost": 10,
                    },
                    {
                        "rationName": "Coomb Milkers",
                        "quantity": 500,
                        "drymatterQuantity": 450,
                        "cost": 50,
                    },
                    {
                        "rationName": "Cwrt Milkers",
                        "quantity": 200,
                        "drymatterQuantity": 180,
                        "cost": 20,
                    },
                    {
                        "rationName": "Coomb Pregnant Heifers",
                        "quantity": 25,
                        "drymatterQuantity": 20,
                        "cost": 4,
                    },
                ],
            }
        ]
    )
    by_farm = {row["farm"]: row for row in rows}
    assert by_farm["GAD"]["as_fed_kg"] == 125
    assert by_farm["GAD"]["dm_kg"] == 110
    assert by_farm["GAD"]["cost"] == 14
    assert by_farm["CM"]["as_fed_kg"] == 200
    assert by_farm["CM"]["dm_kg"] == 180
    assert by_farm["CM"]["cost"] == 20


def test_parse_ingredient_spends_by_farm_excludes_forage_type() -> None:
    rows = _parse_ingredient_spends_by_farm(
        [
            {
                "ingredientName": "Maize Silage",
                "ingredientTypeId": 1,
                "quantity": 400000,
                "rations": [
                    {
                        "rationName": "Coomb Milker Premix",
                        "quantity": 400000,
                        "drymatterQuantity": 140000,
                        "cost": 1,
                    }
                ],
            },
            {
                "ingredientName": "Rape Meal",
                "ingredientTypeId": 2,
                "quantity": 522593,
                "rations": [
                    {
                        "rationName": "Coomb Milker Premix",
                        "quantity": 522593,
                        "drymatterQuantity": 470000,
                        "cost": 150000,
                    }
                ],
            },
            {
                "ingredientName": "Grass Silage",
                "ingredientTypeName": "Forage",
                "quantity": 1000,
                "rations": [
                    {
                        "rationName": "Coomb Milker Premix",
                        "quantity": 1000,
                        "drymatterQuantity": 300,
                        "cost": 1,
                    }
                ],
            },
        ]
    )
    assert [row["ingredient_name"] for row in rows] == ["Rape Meal"]
    assert rows[0]["as_fed_kg"] == 522593


def test_parse_uses_saved_ration_lookup_not_defaults() -> None:
    rows = _parse_ingredient_spends_by_farm(
        [
            {
                "ingredientName": "Rape Meal",
                "ingredientTypeId": 2,
                "rations": [
                    {
                        "rationName": "Coomb Milker Premix",
                        "quantity": 100,
                        "drymatterQuantity": 90,
                        "cost": 10,
                    },
                    {
                        "rationName": "Custom GAD Mix",
                        "quantity": 50,
                        "drymatterQuantity": 45,
                        "cost": 5,
                    },
                ],
            }
        ],
        ration_lookup={"custom gad mix": "GAD"},
    )
    assert len(rows) == 1
    assert rows[0]["as_fed_kg"] == 50


def test_ration_assignments_seed_and_reassign(db: Session) -> None:
    seed_ration_assignments_if_empty(db)
    assert "Coomb Milker Premix" in assigned_ration_names(db, "GAD")
    assert ration_farm_lookup(db)["cwrt milkers"] == "CM"

    save_ration_assignment(db, ration_name="Coomb Milker Premix", farm="")
    save_ration_assignment(db, ration_name="Custom Mix", farm="GAD")
    gad = assigned_ration_names(db, "GAD")
    assert "Coomb Milker Premix" not in gad
    assert "Custom Mix" in gad
    assert "coomb milker premix" not in ration_farm_lookup(db)


def test_parse_respects_saved_ingredient_inclusion() -> None:
    payload = [
        {
            "ingredientName": "Maize Silage",
            "ingredientTypeId": 1,
            "rations": [
                {
                    "rationName": "Coomb Milker Premix",
                    "quantity": 400,
                    "drymatterQuantity": 140,
                    "cost": 1,
                }
            ],
        },
        {
            "ingredientName": "Rape Meal",
            "ingredientTypeId": 2,
            "rations": [
                {
                    "rationName": "Coomb Milker Premix",
                    "quantity": 100,
                    "drymatterQuantity": 90,
                    "cost": 10,
                }
            ],
        },
    ]
    included = _parse_ingredient_spends_by_farm(
        payload,
        ingredient_inclusion={"maize silage": True, "rape meal": False},
    )
    assert [row["ingredient_name"] for row in included] == ["Maize Silage"]
    defaulted = _parse_ingredient_spends_by_farm(payload)
    assert [row["ingredient_name"] for row in defaulted] == ["Rape Meal"]


def test_ingredient_inclusion_lookup_round_trip(db: Session) -> None:
    assert ingredient_inclusion_lookup(db) is None
    save_ingredient_assignment(db, ingredient_name="Maize Silage", included=False)
    save_ingredient_assignment(db, ingredient_name="Rape Meal", included=True)
    lookup = ingredient_inclusion_lookup(db)
    assert lookup is not None
    assert lookup["maize silage"] is False
    assert lookup["rape meal"] is True


def test_usage_report_hides_excluded_stored_ingredients(db: Session) -> None:
    db.add(
        FeedUsageRecord(
            period_start=AUG_START,
            period_end=AUG_END,
            farm="GAD",
            ingredient_name="Maize Silage",
            as_fed_kg=400,
            dm_kg=140,
            cost=1,
        )
    )
    db.add(
        FeedUsageRecord(
            period_start=AUG_START,
            period_end=AUG_END,
            farm="GAD",
            ingredient_name="Rape Meal",
            as_fed_kg=100,
            dm_kg=90,
            cost=10,
        )
    )
    db.commit()
    save_ingredient_assignment(db, ingredient_name="Maize Silage", included=False)
    save_ingredient_assignment(db, ingredient_name="Rape Meal", included=True)
    report = get_usage_report(db, month=AUG_START, farm="GAD")
    assert [row["ingredient_name"] for row in report["ingredients"]] == ["Rape Meal"]
    assert report["totals"]["as_fed_kg"] == 100
    assert report["row_count"] == 1


def test_usage_xlsx_has_borders_and_fitted_columns() -> None:
    report = build_usage_report(
        [
            {
                "ingredient_name": "Megalac / Protected Fat",
                "as_fed_kg": 12_345,
                "dm_kg": 1,
                "cost": 1,
            }
        ],
        period_start=AUG_START,
        period_end=AUG_END,
        farm="GAD",
    )
    workbook = load_workbook(BytesIO(build_usage_xlsx(report)))
    sheet = workbook.active
    assert sheet.title == "Green Acre Dairy"
    assert [cell.value for cell in sheet[1]] == ["Ingredient", "As Fed (kg)", "As Fed (MT)"]
    assert sheet["A2"].value == "Megalac / Protected Fat"
    assert sheet["B2"].value == 12345
    assert sheet["C2"].value == 12.345
    assert sheet["A3"].value == "Total"
    assert sheet["B3"].value == 12345
    assert sheet["C3"].value == 12.345
    assert sheet["B2"].number_format == "#,##0"
    assert sheet["C2"].number_format == "#,##0.0000"
    for row in sheet.iter_rows(min_row=1, max_row=3, min_col=1, max_col=3):
        for cell in row:
            assert cell.border.left.style == "thin"
            assert cell.border.right.style == "thin"
            assert cell.border.top.style == "thin"
            assert cell.border.bottom.style == "thin"
    assert sheet.column_dimensions["A"].width >= len("Megalac / Protected Fat")
    assert sheet.column_dimensions["B"].width >= len("As Fed (kg)")
    assert sheet.column_dimensions["C"].width >= len("As Fed (MT)")

